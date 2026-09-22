"""Executable audit: exercises every endpoint end to end."""
import io, sys, zipfile
from pathlib import Path

from fastapi.testclient import TestClient
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw

sys.path.insert(0, str(Path(__file__).parent))
from main import app  # noqa: E402

FIX = Path("/tmp/fixtures"); FIX.mkdir(exist_ok=True)
client = TestClient(app)
results = []


def build_fixtures():
    img = Image.new("RGB", (1400, 500), "white")
    d = ImageDraw.Draw(img)
    d.text((40, 200), "ILoveWord OCR Test Page", fill="black")
    img.save(FIX / "scan.png")
    big = Image.new("RGB", (3000, 2000), "steelblue")
    big.save(FIX / "big.jpg", quality=95)

    doc = Document()
    doc.add_heading("Quarterly Report", 0)
    doc.add_paragraph("Confidential draft for internal review only.")
    doc.add_paragraph("هذا النص باللغة العربية لاختبار اتجاه الكتابة.")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Region"; t.cell(0, 1).text = "Revenue"
    t.cell(1, 0).text = "EMEA"; t.cell(1, 1).text = "42"
    doc.add_picture(str(FIX / "big.jpg"), width=Inches(4))
    doc.add_page_break()
    doc.add_paragraph("Second page begins here. Replace me please.")
    doc.add_page_break()
    doc.add_paragraph("Third page content.")
    doc.core_properties.author = "Jane Doe"
    doc.save(FIX / "report.docx")

    d2 = Document(); d2.add_paragraph("Appendix A content."); d2.save(FIX / "appendix.docx")
    (FIX / "page.html").write_text("<h1>Title</h1><p>Body <b>text</b>.</p>", encoding="utf-8")
    (FIX / "notes.md").write_text("# Heading\n\n- one\n- two\n\n**bold**\n", encoding="utf-8")
    (FIX / "plain.txt").write_text("line one\nسطر عربي\n", encoding="utf-8")
    (FIX / "fake.docx").write_bytes(b"not a zip at all")


def check(name, resp, *, expect=200, kind=None, min_bytes=200):
    ok = resp.status_code == expect
    detail = ""
    if ok and expect == 200 and kind:
        body = resp.content
        if len(body) < min_bytes:
            ok, detail = False, f"only {len(body)} bytes"
        elif kind == "zip" and not body.startswith(b"PK"):
            ok, detail = False, "not a zip"
        elif kind == "docx":
            if not body.startswith(b"PK"):
                ok, detail = False, "not a docx"
            else:
                try:
                    Document(io.BytesIO(body))
                except Exception as e:
                    ok, detail = False, f"unopenable: {e}"
        elif kind == "pdf" and not body.startswith(b"%PDF"):
            ok, detail = False, "not a pdf"
    if not ok and not detail:
        detail = f"got {resp.status_code}: {resp.text[:160]}"
    results.append((ok, name, detail))
    print(("PASS " if ok else "FAIL ") + name + (f"  [{detail}]" if detail else ""))
    return resp


def f(name, field="file"):
    p = FIX / name
    return {field: (name, p.read_bytes())}


def main():
    build_fixtures()

    r = client.get("/health"); check("health", r)
    print("   engines:", r.json()["engines"])

    # --- tools -------------------------------------------------------------
    check("merge-docx", client.post("/api/tools/merge-docx", files=[
        ("files", ("report.docx", (FIX / "report.docx").read_bytes())),
        ("files", ("appendix.docx", (FIX / "appendix.docx").read_bytes())),
    ]), kind="docx")

    check("split-docx (pages)", client.post("/api/tools/split-docx",
          files=f("report.docx"), data={"mode": "pages"}), kind="zip")
    check("split-docx (count)", client.post("/api/tools/split-docx",
          files=f("report.docx"), data={"mode": "count", "every": "2"}), kind="zip")
    check("compress-docx", client.post("/api/tools/compress-docx",
          files=f("report.docx")), kind="docx")
    check("rtl-fixer", client.post("/api/tools/rtl-fixer", files=f("report.docx")), kind="docx")
    check("watermark-word", client.post("/api/tools/watermark-word", files=f("report.docx"),
          data={"text": "سري - CONFIDENTIAL"}), kind="docx")
    check("remove-blank-pages", client.post("/api/tools/remove-blank-pages",
          files=f("report.docx")), kind="docx")
    check("find-replace", client.post("/api/tools/find-replace", files=f("report.docx"),
          data={"search": "Replace me", "replace": "REPLACED"}), kind="docx")
    check("clean-formatting", client.post("/api/tools/clean-formatting",
          files=f("report.docx")), kind="docx")
    check("remove-metadata", client.post("/api/tools/remove-metadata",
          files=f("report.docx")), kind="docx")
    check("extract-text", client.post("/api/tools/extract-text", files=f("report.docx")),
          min_bytes=20, kind="txt")
    check("extract-images", client.post("/api/tools/extract-images", files=f("report.docx")),
          kind="zip")
    r = check("word-stats", client.post("/api/tools/word-stats", files=f("report.docx")))
    print("   stats:", r.json()["stats"])

    prot = check("protect-word", client.post("/api/tools/protect-word", files=f("report.docx"),
                 data={"password": "s3cret!"}), min_bytes=500)
    (FIX / "locked.docx").write_bytes(prot.content)
    check("unlock-word (correct pw)", client.post("/api/tools/unlock-word",
          files=f("locked.docx"), data={"password": "s3cret!"}), kind="docx")
    check("unlock-word (wrong pw)", client.post("/api/tools/unlock-word",
          files=f("locked.docx"), data={"password": "nope"}), expect=401)
    check("unlock-word (no pw)", client.post("/api/tools/unlock-word",
          files=f("locked.docx")), expect=401)

    # --- conversions -------------------------------------------------------
    pdf = check("word-to-pdf", client.post("/api/convert/word-to-pdf", files=f("report.docx")),
                kind="pdf")
    if pdf.status_code == 200:
        (FIX / "report.pdf").write_bytes(pdf.content)
        check("pdf-to-docx", client.post("/api/convert/pdf-to-docx", files=f("report.pdf")),
              kind="docx")
    check("word-to-jpg", client.post("/api/convert/word-to-jpg", files=f("report.docx"),
          data={"dpi": "100"}), kind="zip")
    check("image-to-docx (OCR)", client.post("/api/convert/image-to-docx", files=f("scan.png")),
          kind="docx")
    check("word-to-html", client.post("/api/convert/word-to-html", files=f("report.docx")),
          min_bytes=100)
    check("html-to-word", client.post("/api/convert/html-to-word", files=f("page.html")),
          kind="docx")
    check("markdown-to-word", client.post("/api/convert/markdown-to-word", files=f("notes.md")),
          kind="docx")
    check("text-to-word", client.post("/api/convert/text-to-word", files=f("plain.txt")),
          kind="docx")

    # --- negative paths: the 422 / 500 hunt --------------------------------
    check("missing file field", client.post("/api/tools/extract-text"), expect=422)
    check("wrong field name", client.post("/api/tools/extract-text",
          files={"document": ("a.docx", (FIX / "report.docx").read_bytes())}), expect=422)
    check("merge with single file", client.post("/api/tools/merge-docx",
          files=[("files", ("a.docx", (FIX / "report.docx").read_bytes()))]), expect=400)
    check("wrong extension", client.post("/api/tools/extract-text", files=f("plain.txt")),
          expect=415)
    check("corrupt docx", client.post("/api/tools/extract-text", files=f("fake.docx")),
          expect=400)
    check("empty file", client.post("/api/tools/extract-text",
          files={"file": ("empty.docx", b"")}), expect=400)
    check("watermark missing text", client.post("/api/tools/watermark-word",
          files=f("report.docx")), expect=422)
    check("find-replace no match", client.post("/api/tools/find-replace", files=f("report.docx"),
          data={"search": "zzzzz", "replace": "x"}), expect=400)
    check("encrypted file to normal tool", client.post("/api/tools/extract-text",
          files=f("locked.docx")), expect=400)
    check("path traversal filename", client.post("/api/tools/extract-text",
          files={"file": ("../../etc/passwd.docx", (FIX / "report.docx").read_bytes())}))

    # error body shape
    r = client.post("/api/tools/extract-text", files=f("plain.txt"))
    body = r.json()
    check("error shape {detail, code}", r, expect=415)
    assert "detail" in body and "code" in body, body

    # verify content of a couple of outputs
    r = client.post("/api/tools/find-replace", files=f("report.docx"),
                    data={"search": "Replace me", "replace": "REPLACED"})
    txt = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    results.append(("REPLACED" in txt, "find-replace actually replaced", ""))
    print(("PASS " if "REPLACED" in txt else "FAIL ") + "find-replace actually replaced")

    r = client.post("/api/tools/split-docx", files=f("report.docx"), data={"mode": "pages"})
    n = len(zipfile.ZipFile(io.BytesIO(r.content)).namelist())
    results.append((n >= 3, f"split produced {n} parts", ""))
    print(("PASS " if n >= 3 else "FAIL ") + f"split produced {n} parts")

    r = client.post("/api/tools/remove-metadata", files=f("report.docx"))
    author = Document(io.BytesIO(r.content)).core_properties.author
    results.append((not author, "metadata stripped", f"author={author!r}"))
    print(("PASS " if not author else "FAIL ") + "metadata stripped")

    failed = [r for r in results if not r[0]]
    print(f"\n{len(results) - len(failed)}/{len(results)} passed")
    for _, name, detail in failed:
        print("  FAILED:", name, detail)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
